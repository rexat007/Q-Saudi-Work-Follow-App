"""Operator-only deterministic XLSX -> Google blueprint authoring. Requires openpyxl.
Never imported by runtime. Usage: python3 ... --xlsx source.xlsx --out /tmp/workbook.json
The approved script sources are versioned separately; do not execute the source XLSX.
"""
import argparse,datetime,hashlib,json,re,colorsys
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import range_boundaries,get_column_letter
from xml.etree import ElementTree as ET

def extract(path):
 w=openpyxl.load_workbook(path,data_only=False);report={'sourceSha256':hashlib.sha256(open(path,'rb').read()).hexdigest(),'sourceFormulaCells':0,'removedCells':[],'normalizedFormulaCells':0,'classification':{}}
 theme=[]
 if w.loaded_theme:
  t=ET.fromstring(w.loaded_theme);ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
  scheme=t.find('.//a:clrScheme',ns)
  for c in scheme:theme.append(next(iter(c)).get('val') if next(iter(c)).get('val') not in ('window','windowText') else next(iter(c)).get('lastClr'))
 def color(c):
  if c is None:return None
  v=c.rgb if c.type=='rgb' else theme[c.theme] if c.type=='theme' and c.theme<len(theme) else None
  if not v:return None
  v=v[-6:];rgb=[int(v[i:i+2],16)/255 for i in (0,2,4)];t=c.tint or 0
  rgb=[x*(1+t) if t<0 else x*(1-t)+t for x in rgb]
  return dict(zip(['red','green','blue'],rgb))
 def fmt(c):
  f={};font=c.font
  tf={k:v for k,v in {'bold':bool(font.bold),'italic':bool(font.italic),'strikethrough':bool(font.strike),'fontFamily':font.name,'fontSize':int(font.sz) if font.sz else None,'foregroundColor':color(font.color)}.items() if v is not None}
  if font.underline:tf['underline']=True
  if tf:f['textFormat']=tf
  if c.fill.patternType=='solid' and color(c.fill.fgColor):f['backgroundColor']=color(c.fill.fgColor)
  align=c.alignment
  if align.horizontal in ['left','center','right']:f['horizontalAlignment']=align.horizontal.upper()
  if align.vertical in ['top','center','bottom']:f['verticalAlignment']='MIDDLE' if align.vertical=='center' else align.vertical.upper()
  if align.wrap_text:f['wrapStrategy']='WRAP'
  if align.textRotation and align.textRotation<=180:f['textRotation']={'angle':align.textRotation if align.textRotation<=90 else 90-align.textRotation}
  borders={};styles={'thin':'SOLID','medium':'SOLID_MEDIUM','thick':'SOLID_THICK','double':'DOUBLE','dotted':'DOTTED','dashed':'DASHED','hair':'DOTTED','mediumDashed':'DASHED','dashDot':'DASHED','mediumDashDot':'DASHED','dashDotDot':'DASHED','mediumDashDotDot':'DASHED','slantDashDot':'DASHED'}
  for side in ['top','bottom','left','right']:
   b=getattr(c.border,side)
   if b and b.style:
    if b.style not in styles:raise ValueError('Unmapped border '+b.style)
    borders[side]={'style':styles[b.style],**({'color':color(b.color)} if color(b.color) else {})}
  if borders:f['borders']=borders
  nf=c.number_format
  if nf and nf!='General':
   low=re.sub(r'"[^\"]*"','',nf).lower();typ='TEXT' if nf=='@' else 'PERCENT' if '%' in low else 'DATE_TIME' if re.search('[yd]',low) and re.search('[hs]',low) else 'DATE' if re.search('[yd]',low) else 'TIME' if re.search('[hs]',low) else 'NUMBER'
   f['numberFormat']={'type':typ,'pattern':nf}
  return f
 def grid(sid,rng):
  a,b,c,d=range_boundaries(str(rng));return {'sheetId':sid,'startRowIndex':(b or 1)-1,'endRowIndex':min(d,nr) if d else nr,'startColumnIndex':(a or 1)-1,'endColumnIndex':min(c,nc) if c else nc}
 def formula(value):
  v=value if isinstance(value,str) else value.text
  return v.replace('_xlfn.','').replace('_xlws.','')
 def normalize(v):
  # All operational lookup/aggregate ranges follow source growth; D11 is handled separately.
  return re.sub(r"((?:'Loading Register'|'Offload Data')!\$?[A-Z]{1,3}\$?2:\$?[A-Z]{1,3})\$?(?:102|298|979|1000|10189)(?!\d)",r'\1',v)
 sheets=[]
 for index,s in enumerate(w):
  sid=100+index;title=s.title
  if s._charts or s.protection.sheet or len(w.defined_names):raise ValueError('New unsupported structural feature requires explicit capture review')
  nr=298 if index<2 else max(s.max_row,40);nc=max(s.max_column,12 if title=='Carrier Report Viewer' else 1)
  cells={};sourceformula=0;keepcount=0;removed=0
  for row in s:
   for c in row:
    if index<2 and c.row>nr:
     if c.value is not None:raise ValueError('Cannot trim nonempty operational cells')
     continue
    if not c.has_style and c.value is None:continue
    data={};f=fmt(c)
    if f:data['userEnteredFormat']=f
    v=c.value
    if c.data_type=='f':
     sourceformula+=1;v=normalize(formula(v))
     # Empty guards count operator inputs only, not computed zero-valued cells.
     if title=='Loading Register' and c.coordinate=='U2':v=v.replace('COUNTA(A2:C2,E2:F2,O2:T2)=0','COUNTA(A2:C2,E2:F2,O2:S2)=0')
     if title=='Offload Data' and c.coordinate=='M2':v=v.replace('COUNTA(B2:L2)=0','COUNTA(B2:G2,I2:L2)=0')
     if title=='Dashboard' and c.coordinate=='D11':v="=ROWS('Loading Register'!$A$2:$A$298)-COUNTIF('Loading Register'!$A$2:$A$298,\"<>\")"
     data['userEnteredValue']={'formulaValue':v};keepcount+=1
    elif v is not None:
     keep=True
     if index<2:keep=c.row==1
     elif title=='Offload Operations Audit':keep=c.row in (1,3,5) or c.coordinate in ['A2','D2','G2','J2','K2']
     elif title=='Carrier Daily Report':keep=c.row in (1,3,5) or c.coordinate in ['A2','B2','D2','G2','J2']
     elif title=='Carrier Report Viewer':keep=c.row in (1,3,4,5,6,7,10) and c.coordinate not in ['B3','B4','E3']
     elif title=='Daily Summary':keep=c.coordinate!='B3'
     elif title=='Dashboard':keep=not ((c.column in (21,23) and 5<=c.row<=12) or (c.row==27 and c.column in (2,4)))
     if keep:
      if isinstance(v,datetime.datetime):raise ValueError('Unclassified historical date '+title+'!'+c.coordinate)
      key='boolValue' if isinstance(v,bool) else 'numberValue' if isinstance(v,(int,float)) else 'stringValue';data['userEnteredValue']={key:v};keepcount+=1
     else:removed+=1;report['removedCells'].append(title+'!'+c.coordinate)
    cells[(c.row,c.column)]=data
  report['sourceFormulaCells']+=sourceformula
  # Prepare the 297 operational rows intentionally using the supplied row-2 seeds.
  if index<2:
   for col in range(1,nc+1):
    seed=cells.get((2,col),{});fv=seed.get('userEnteredValue',{}).get('formulaValue')
    for row in range(3,299):
     if fv:cells.setdefault((row,col),{})['userEnteredValue']={'formulaValue':Translator(fv,origin=get_column_letter(col)+'2').translate_formula(get_column_letter(col)+str(row))}
     if seed.get('userEnteredFormat') and not cells.get((row,col),{}).get('userEnteredFormat'):cells.setdefault((row,col),{})['userEnteredFormat']=seed['userEnteredFormat']
  def setcell(addr,v):
   c=s[addr];key='formulaValue' if isinstance(v,str) and v.startswith('=') else 'numberValue' if isinstance(v,(int,float)) else 'stringValue';cells.setdefault((c.row,c.column),{})['userEnteredValue']={key:v}
  if title=='Offload Operations Audit':
   for a in ['E2','H2']:setcell(a,0)
  if title=='Carrier Daily Report':
   for a in ['H2','K2']:setcell(a,0)
  if title=='Carrier Report Viewer':
   for c in range(1,9):setcell(get_column_letter(c)+'8',0)
   for c,v in enumerate(['Carrier / الناقل','Variance >1t','Unverified offloads','Arrival deficit EOD','Offload deficit EOD'],1):setcell(get_column_letter(c)+'12',v)
   setcell('E3','No records');setcell('A6','No daily ledger data. Refresh Carrier Reports after source edits.')
  if title=='Dashboard':
   for row in range(5,10):
    setcell('U'+str(row),0);setcell('W'+str(row),1 if row==9 else 0)
   setcell('A11','Net Weight Variance (tonnes)')
  for dv in s.data_validations.dataValidation:
   condition=None
   if dv.type=='list':
    vals=(dv.formula1 or '').strip('"').split(',')
    if title=='Carrier Report Viewer':vals=['All']
    condition={'type':'ONE_OF_LIST','values':[{'userEnteredValue':v} for v in vals]}
   elif dv.type=='decimal':condition={'type':{'greaterThan':'NUMBER_GREATER','greaterThanOrEqual':'NUMBER_GREATER_THAN_EQ','between':'NUMBER_BETWEEN'}.get(dv.operator,'NUMBER_GREATER'),'values':[{'userEnteredValue':dv.formula1 or '0'}]+([{'userEnteredValue':dv.formula2}] if dv.formula2 else [])}
   elif dv.type=='custom':
    custom=normalize(dv.formula1 if (dv.formula1 or '').startswith('=') else '='+dv.formula1)
    # Excel's CELL("format") result codes are not a Google Sheets API contract.
    # Date serials emitted by the canonical app and parseable date text remain valid.
    if 'CELL("format", B2)' in custom:custom='=OR(ISNUMBER(B2),NOT(ISERROR(DATEVALUE(B2))))'
    condition={'type':'CUSTOM_FORMULA','values':[{'userEnteredValue':custom}]}
   else:raise ValueError('Unmapped validation '+str(dv.type))
   for rng in dv.sqref.ranges:
    for row in range(rng.min_row,min(rng.max_row,nr)+1):
     for col in range(rng.min_col,min(rng.max_col,nc)+1):cells.setdefault((row,col),{})['dataValidation']={'condition':condition,'strict':dv.error is not None or bool(dv.showErrorMessage),'showCustomUi':True}
  if title=='Offload Data':
   for r in range(2,nr+1):cells.setdefault((r,10),{})['dataValidation']={'condition':{'type':'ONE_OF_LIST','values':[{'userEnteredValue':'Ali Hadi'},{'userEnteredValue':'Abbass Abdelhalim'}]},'strict':True,'showCustomUi':True}
  cfs=[]
  for cf,rules in s.conditional_formatting._cf_rules.items():
   for rule in rules:
    if rule.type not in ('expression','containsText') or not rule.formula:raise ValueError('Unmapped conditional rule '+rule.type)
    obj={'ranges':[grid(sid,r) for r in cf.sqref.ranges],'booleanRule':{'condition':{'type':'CUSTOM_FORMULA','values':[{'userEnteredValue':'='+str(rule.formula[0]).lstrip('=')}]},'format':{}}}
    if rule.dxf:
     if rule.dxf.fill and color(rule.dxf.fill.fgColor):obj['booleanRule']['format']['backgroundColor']=color(rule.dxf.fill.fgColor)
     if rule.dxf.font:obj['booleanRule']['format']['textFormat']={k:v for k,v in {'bold':rule.dxf.font.b,'foregroundColor':color(rule.dxf.font.color)}.items() if v is not None}
    cfs.append(obj)
  # XLSX table wrappers are export metadata, not identity/structured-reference authority.
  # Reproduce their actual stripe colours as Sheets banding and their existing filters.
  bandings=[];basic_filter=None
  for table in s.tables.values():
   table_range=grid(sid,table.ref)
   if index<2:table_range['endRowIndex']=nr
   style=next((t for t in w._table_styles.tableStyle if t.name==table.tableStyleInfo.name),None)
   if not style:raise ValueError('Unmapped table style '+table.name)
   props={}
   for el in style.tableStyleElement:
    role={'headerRow':'headerColor','firstRowStripe':'firstBandColor','secondRowStripe':'secondBandColor'}.get(el.type)
    if role:
     d=w._differential_styles.styles[el.dxfId]
     if d.fill and color(d.fill.fgColor):props[role]=color(d.fill.fgColor)
   whole=next((w._differential_styles.styles[e.dxfId] for e in style.tableStyleElement if e.type=='wholeTable'),None)
   if whole and whole.border:
    tb={}
    for side in ['top','bottom','left','right']:
     edge=getattr(whole.border,side)
     if edge and edge.style:
      if edge.style!='thin':raise ValueError('Unmapped table border style')
      tb[side]={'style':'SOLID',**({'color':color(edge.color)} if color(edge.color) else {})}
    for r in range(table_range['startRowIndex']+1,table_range['endRowIndex']+1):
     for c in range(table_range['startColumnIndex']+1,table_range['endColumnIndex']+1):
      f=cells.setdefault((r,c),{}).setdefault('userEnteredFormat',{});f['borders']={**tb,**f.get('borders',{})}
   if table.tableStyleInfo.showRowStripes and props:bandings.append({'bandedRangeId':20000+sid,'range':table_range,'rowProperties':props})
   if table.autoFilter and table.autoFilter.ref:basic_filter={'range':table_range}
  if s.auto_filter.ref:basic_filter={'range':grid(sid,s.auto_filter.ref)}
  maxr=max([r for r,c in cells],default=1);rows=[];rowends={}
  for r,c in cells:rowends[r]=max(rowends.get(r,0),c)
  for r in range(1,maxr+1):
   end=rowends.get(r,0);rows.append({'values':[cells.get((r,c),{}) for c in range(1,end+1)]})
  colmeta=[{} for _ in range(nc)];rowmeta=[{} for _ in range(nr)]
  for key,d in s.column_dimensions.items():
   for c in range(d.min or openpyxl.utils.column_index_from_string(key),min(d.max or openpyxl.utils.column_index_from_string(key),nc)+1):colmeta[c-1]={'pixelSize':round((d.width or 13)*7+5),'hiddenByUser':bool(d.hidden)}
  for r,d in s.row_dimensions.items():
   if r<=nr:rowmeta[r-1]={'hiddenByUser':bool(d.hidden),**({'pixelSize':round(d.height*96/72)} if d.height else {})}
  frozen={'frozenRowCount':s[s.freeze_panes].row-1,'frozenColumnCount':s[s.freeze_panes].column-1} if s.freeze_panes else {}
  sheets.append({'properties':{'sheetId':sid,'title':title,'index':index,'gridProperties':{'rowCount':nr,'columnCount':nc,**frozen}},'data':[{'startRow':0,'startColumn':0,'rowData':rows,'columnMetadata':colmeta,'rowMetadata':rowmeta}],**({'merges':[grid(sid,r) for r in s.merged_cells.ranges]} if s.merged_cells.ranges else {}),**({'conditionalFormats':cfs} if cfs else {}),**({'bandedRanges':bandings} if bandings else {}),**({'basicFilter':basic_filter} if basic_filter else {})})
  report['normalizedFormulaCells']+=sum('formulaValue' in d.get('userEnteredValue',{}) for d in cells.values());report['classification'][title]={'sourceFormulaCells':sourceformula,'removedLiteralCells':removed,'retainedSourceCells':keepcount,'rows':nr,'columns':nc}
 return {'schemaVersion':1,'version':'SABIL-HAIL-V1','approved':True,'properties':{'locale':'en_US','timeZone':'Asia/Riyadh'},'sheets':sheets,'namedRanges':[],'capture':{'complete':True,'unaccounted':0,'operationalRows':0,'sourceReferenceIds':[]}},report
def compact(blueprint):
 styles=[];styleids={}
 def rle(items):
  out=[]
  for item in items:
   if out and out[-1][1]==item:out[-1][0]+=1
   else:out.append([1,item])
  return out
 for sheet in blueprint['sheets']:
  for data in sheet['data']:
   for row in data['rowData']:
    for cell in row['values']:
     if 'userEnteredFormat' in cell:
      style=cell.pop('userEnteredFormat');key=json.dumps(style,sort_keys=True)
      if key not in styleids:styleids[key]=len(styles);styles.append(style)
      cell['styleRef']=styleids[key]
   for field in ['rowData','rowMetadata','columnMetadata']:data[field+'Rle']=rle(data.pop(field))
 blueprint['formatPool']=styles
 return blueprint

if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--xlsx',required=True);ap.add_argument('--out',required=True);args=ap.parse_args();b,r=extract(args.xlsx);b=compact(b)
 with open(args.out,'w') as f:json.dump(b,f,ensure_ascii=False,sort_keys=True,separators=(',',':'));f.write('\n')
 with open(args.out+'.report.json','w') as f:json.dump(r,f,indent=2,ensure_ascii=False)
